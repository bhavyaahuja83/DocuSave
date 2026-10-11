from pathlib import Path

from docx import Document
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {".docx", ".xlsx"}


def extract_text(path: str | Path) -> str:
    """Return readable paragraph/table or worksheet text for preview and diffs."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        document = Document(path)
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table_index, table in enumerate(document.tables, start=1):
            lines.append(f"[Table {table_index}]")
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(lines)

    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            lines = []
            for sheet in workbook.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    while values and not values[-1]:
                        values.pop()
                    if values:
                        lines.append(" | ".join(values))
            return "\n".join(lines)
        finally:
            workbook.close()

    raise ValueError(f"Unsupported document type: {suffix or '(no extension)'}")


def replace_text(path: str | Path, content: str) -> None:
    """Update readable text while requiring the document's existing structure."""
    path = Path(path)
    lines = content.splitlines()

    if path.suffix.lower() == ".docx":
        document = Document(path)
        expected_count = len(document.paragraphs) + sum(
            1 + len(table.rows) for table in document.tables
        )
        if len(lines) != expected_count:
            raise ValueError("Keep the existing paragraph and table-row structure while editing.")

        cursor = 0
        for paragraph in document.paragraphs:
            _replace_paragraph_text(paragraph, lines[cursor])
            cursor += 1
        for table_index, table in enumerate(document.tables, start=1):
            if lines[cursor] != f"[Table {table_index}]":
                raise ValueError("Table labels and order cannot be changed in the editor.")
            cursor += 1
            for row in table.rows:
                values = lines[cursor].split(" | ")
                if len(values) != len(row.cells):
                    raise ValueError("Keep the same number of cells in each table row.")
                for cell, value in zip(row.cells, values):
                    _replace_paragraph_text(cell.paragraphs[0], value)
                cursor += 1
        document.save(path)
        return

    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, data_only=False)
        try:
            cursor = 0
            updates = []
            for sheet in workbook.worksheets:
                if cursor >= len(lines) or lines[cursor] != f"[Sheet: {sheet.title}]":
                    raise ValueError("Worksheet names and order cannot be changed in the editor.")
                cursor += 1
                for row in sheet.iter_rows():
                    if not any(cell.value is not None for cell in row):
                        continue
                    if cursor >= len(lines):
                        raise ValueError("Keep the existing worksheet row structure while editing.")
                    values = lines[cursor].split(" | ")
                    existing_width = len(row)
                    if len(values) > existing_width:
                        raise ValueError("A worksheet row cannot contain more cells than it has now.")
                    updates.append((row, values))
                    cursor += 1
            if cursor != len(lines):
                raise ValueError("Keep worksheet names, rows, and order unchanged while editing.")

            for row, values in updates:
                for index, cell in enumerate(row):
                    value = values[index] if index < len(values) else ""
                    original = "" if cell.value is None else str(cell.value)
                    if value != original:
                        cell.value = _parse_cell_value(value)
            workbook.save(path)
        finally:
            workbook.close()
        return

    raise ValueError(f"Unsupported document type: {path.suffix or '(no extension)'}")


def _parse_cell_value(value: str) -> str | int | float:
    if value.startswith("="):
        return value
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def _replace_paragraph_text(paragraph, value: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(value)
        return
    paragraph.runs[0].text = value
    for run in paragraph.runs[1:]:
        run.text = ""
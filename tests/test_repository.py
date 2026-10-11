import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from docugit.extractors import extract_text, replace_text
from docugit.repository import DocumentRepository


class DocumentRepositoryTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)
        self.repository = DocumentRepository(self.root / "repository")
        self.source = self.root / "source.docx"
        self._write_docx(self.source, "First draft")

    def tearDown(self):
        self.temporary_directory.cleanup()

    @staticmethod
    def _write_docx(path: Path, text: str) -> None:
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    def test_add_update_diff_and_restore(self):
        document = self.repository.add_document(self.source, "Initial import", "Test User")
        tracked_copy = self.repository.current_path(document["id"])
        self.assertEqual(tracked_copy.read_bytes(), self.source.read_bytes())
        self.assertEqual(self.repository.list_versions(document["id"])[0]["author"], "Test User")

        updated = self.root / "updated.docx"
        self._write_docx(updated, "Second draft")
        version = self.repository.create_version(document["id"], updated, "Revise intro", "Test User")
        self.assertEqual(version["version_number"], 2)
        diff = self.repository.compare_versions(document["id"], 1, 2)
        self.assertIn("-First draft", diff)
        self.assertIn("+Second draft", diff)
        self.assertEqual(tracked_copy.read_bytes(), updated.read_bytes())

        restored = self.repository.restore_version(document["id"], 1, "Test User")
        self.assertEqual(restored["version_number"], 3)
        self.assertEqual(tracked_copy.read_bytes(), self.source.read_bytes())
        self.assertEqual(self.repository.document_text(document["id"]), "First draft")

    def test_duplicate_names_get_distinct_working_copies(self):
        first = self.repository.add_document(self.source, "First", "Test User")
        second = self.repository.add_document(self.source, "Second", "Test User")
        self.assertNotEqual(first["file_name"], second["file_name"])
        self.assertEqual(second["file_name"], "source (2).docx")

    def test_commit_current_snapshots_the_tracked_working_copy(self):
        document = self.repository.add_document(self.source, "Initial import", "Test User")
        working_copy = self.repository.current_path(document["id"])
        self._write_docx(working_copy, "Edited in Word")

        version = self.repository.commit_current(document["id"], "Update text", "Test User")

        self.assertEqual(version["version_number"], 2)
        self.assertIn("Edited in Word", self.repository.compare_versions(document["id"], 1, 2))
        self.assertEqual(self.source.read_bytes(), self.repository.version_path(
            self.repository.get_version(document["id"], 1)
        ).read_bytes())
        with self.assertRaisesRegex(ValueError, "no changes to commit"):
            self.repository.commit_current(document["id"], "Duplicate", "Test User")

    def test_publish_commits_edits_and_archives_the_previous_original(self):
        document = self.repository.add_document(self.source, "Initial import", "Test User")
        original_before = self.source.read_bytes()
        self._write_docx(self.repository.current_path(document["id"]), "Ready to publish")

        result = self.repository.publish_current(document["id"], "Publish edits", "Test User")

        self.assertEqual(result["version"]["version_number"], 2)
        self.assertEqual(result["backup_path"].read_bytes(), original_before)
        self.assertEqual(self.source.read_bytes(), self.repository.current_path(document["id"]).read_bytes())
        self.assertEqual(result["original_path"], self.source.resolve())

        second_publish = self.repository.publish_current(
            document["id"], "No new edits", "Test User"
        )
        self.assertEqual(second_publish["version"]["version_number"], 2)
        self.assertEqual(len(self.repository.list_versions(document["id"])), 2)

    def test_in_app_editor_updates_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Opening")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        document.add_paragraph("Closing")
        path = self.root / "editable.docx"
        document.save(path)

        replace_text(path, "Introduction\nConclusion\n[Table 1]\nItem | 42")

        self.assertEqual(
            extract_text(path), "Introduction\nConclusion\n[Table 1]\nItem | 42"
        )
        with self.assertRaisesRegex(ValueError, "paragraph and table-row structure"):
            replace_text(path, "Only one line")

    def test_in_app_editor_updates_excel_values(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Budget"
        sheet.append(["Item", "Cost"])
        sheet.append(["Travel", 25])
        path = self.root / "editable.xlsx"
        workbook.save(path)

        replace_text(path, "[Sheet: Budget]\nItem | Cost\nTravel | 42")

        self.assertEqual(extract_text(path), "[Sheet: Budget]\nItem | Cost\nTravel | 42")
        self.assertEqual(load_workbook(path, data_only=False)["Budget"]["B2"].value, 42)

    def test_excel_preview_and_extension_validation(self):
        workbook_path = self.root / "sheet.xlsx"
        workbook = Workbook()
        workbook.active.title = "Budget"
        workbook.active.append(["Item", "Cost"])
        workbook.active.append(["Travel", 25])
        workbook.save(workbook_path)
        document = self.repository.add_document(workbook_path, "Budget import", "Test User")
        preview = self.repository.document_text(document["id"])
        self.assertIn("[Sheet: Budget]", preview)
        self.assertIn("Travel | 25", preview)

        with self.assertRaisesRegex(ValueError, "requires a .xlsx"):
            self.repository.create_version(document["id"], self.source, "Wrong type", "Test User")


if __name__ == "__main__":
    unittest.main()